from openai import OpenAI
from dotenv import load_dotenv
from rich.console import Console
from rich.markdown import Markdown
from rich.panel import Panel
from rich.text import Text

from prompt_toolkit import PromptSession
from prompt_toolkit.key_binding import KeyBindings

import html
import json
import logging
import os
import re
import sys
import threading
import time
from pathlib import Path
from typing import Iterable
from pptx import Presentation
from docx import Document
from pypdf import PdfReader
def get_base_dir() -> Path:
    """返回程序所在目录，兼容 PyInstaller 打包后的 exe。"""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


BASE_DIR = get_base_dir()
CURRENT_DIR = Path.cwd()
console = Console()


DEFAULT_CONFIG = {
    "api_base_url": "https://api.deepseek.com",
    "models": ["deepseek-v4-pro", "deepseek-v4", "deepseek-v3"],
    "temperature": 0.7,
    "timeout": 60,
    "retry_times": 2,
    "knowledge_path": "knowledge_test.txt",
    "max_history_messages": 20,
    "knowledge_top_k": 4,
    "logs_dir": "logs",
    "max_upload_chars": 30000,
}


def load_config() -> dict:
    config = DEFAULT_CONFIG.copy()
    for path in (CURRENT_DIR / "config.json", BASE_DIR / "config.json"):
        if not path.exists():
            continue
        try:
            with path.open("r", encoding="utf-8") as f:
                config.update(json.load(f))
            break
        except Exception as exc:
            console.print(f"[yellow]config.json 读取失败，已使用默认配置：{exc}[/yellow]")
    return config


CONFIG = load_config()


def setup_environment() -> None:
    for path in (CURRENT_DIR / ".env", BASE_DIR / ".env"):
        if path.exists():
            load_dotenv(path)
            return
    load_dotenv()


def setup_logging() -> None:
    logs_dir = CURRENT_DIR / CONFIG["logs_dir"]
    logs_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        filename=logs_dir / "assistant.log",
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        encoding="utf-8",
    )


setup_environment()
setup_logging()


def resolve_path(path_text: str) -> Path:
    path_text = path_text.strip().strip('"').strip("'")
    path = Path(path_text)
    if path.is_absolute():
        return path
    return CURRENT_DIR / path


def load_text_file(path: Path, max_chars: int | None = None) -> tuple[str, bool]:
    with path.open("r", encoding="utf-8") as f:
        text = f.read()

    if max_chars is not None and len(text) > max_chars:
        return text[:max_chars], True
    return text, False
def load_ppt_file(path: Path, max_chars: int | None = None) -> tuple[str, bool]:
    """读取 pptx 文本内容。"""
    prs = Presentation(path)

    texts = []

    for slide_index, slide in enumerate(prs.slides, start=1):
        texts.append(f"\n===== 第 {slide_index} 页 =====\n")

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    texts.append(text)

    full_text = "\n".join(texts)

    if max_chars is not None and len(full_text) > max_chars:
        return full_text[:max_chars], True

    return full_text, False
def load_docx_file(path: Path, max_chars: int | None = None) -> tuple[str, bool]:
    """读取 docx 文本内容。"""
    doc = Document(path)

    texts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            texts.append(text)

    full_text = "\n".join(texts)

    if max_chars is not None and len(full_text) > max_chars:
        return full_text[:max_chars], True

    return full_text, False
def load_pdf_file(path: Path, max_chars: int | None = None) -> tuple[str, bool]:
    """读取 PDF 文本内容。"""
    reader = PdfReader(path)

    texts = []

    for page_index, page in enumerate(reader.pages, start=1):
        texts.append(f"\n===== 第 {page_index} 页 =====\n")

        try:
            text = page.extract_text()
            if text:
                texts.append(text.strip())
        except Exception:
            continue

    full_text = "\n".join(texts)

    if max_chars is not None and len(full_text) > max_chars:
        return full_text[:max_chars], True

    return full_text, False
def load_knowledge() -> str:
    candidates = [
        resolve_path(CONFIG["knowledge_path"]),
        BASE_DIR / CONFIG["knowledge_path"],
    ]
    for path in candidates:
        if path.exists():
            try:
                text, _ = load_text_file(path)
                return text
            except Exception as exc:
                console.print(f"[yellow]知识库读取失败：{exc}[/yellow]")

    console.print("[yellow]未找到 knowledge_test.txt，知识库内容将为空。[/yellow]")
    return ""


KNOWLEDGE = load_knowledge()


BASE_SYSTEM_PROMPT = """
你是一名大学学习助手。

回答要求：
1. 解答大学问题，通俗易懂，像老师一样讲解。
2. 多举例，使用 Markdown 格式。
3. 如果有代码，请使用 Markdown 代码块。
4. 如果有数学公式，请使用 LaTeX，例如 $E=mc^2$ 或 $$a^2+b^2=c^2$$。

知识库使用规则：
1. 如果本轮问题检索到了知识库片段，必须优先参考这些内容。
2. 可以在知识库基础上补充解释。
3. 如果没有检索到相关内容，再使用你的通用知识回答。

代码文件输出要求：
当用户明确需要完整的 .py、.cpp、.md 文件时：
1. 必须给出完整文件内容。
2. 必须使用代码块，不要省略 import、main 函数、类定义等必要部分。
3. 不要写“其他部分不变”“这里省略”“同上”。
4. 如果用户指定文件名，请在代码块开头写成：
   ```python filename=main.py
   ```cpp filename=main.cpp
   ```markdown filename=README.md
"""


CHAT_FILE = CURRENT_DIR / "chat_history.json"
HTML_REPLY_FILE = CURRENT_DIR / "latest_reply.html"
MAX_HISTORY_MESSAGES = int(CONFIG["max_history_messages"])


def normalize_words(text: str) -> set[str]:
    return set(re.findall(r"[\w\u4e00-\u9fff]+", text.lower()))


def split_knowledge(knowledge: str) -> list[str]:
    chunks = [chunk.strip() for chunk in re.split(r"\n\s*\n", knowledge) if chunk.strip()]
    return chunks or ([knowledge.strip()] if knowledge.strip() else [])


def retrieve_knowledge(query: str, knowledge: str, top_k: int) -> str:
    """轻量 RAG：按关键词重叠度选出最相关的知识片段。"""
    query_words = normalize_words(query)
    if not query_words:
        return ""

    scored_chunks = []
    for chunk in split_knowledge(knowledge):
        chunk_words = normalize_words(chunk)
        score = len(query_words & chunk_words)
        if score:
            scored_chunks.append((score, len(chunk), chunk))

    scored_chunks.sort(key=lambda item: (-item[0], item[1]))
    return "\n\n".join(chunk for _, _, chunk in scored_chunks[:top_k])


def build_system_prompt(user_input: str) -> str:
    relevant_knowledge = retrieve_knowledge(
        user_input,
        KNOWLEDGE,
        int(CONFIG["knowledge_top_k"]),
    )
    if not relevant_knowledge:
        return BASE_SYSTEM_PROMPT

    return (
        BASE_SYSTEM_PROMPT
        + "\n\n========================\n【本轮相关知识库片段】\n========================\n"
        + relevant_knowledge
    )


def load_chat_history() -> list[dict]:
    if not CHAT_FILE.exists():
        return []
    try:
        with CHAT_FILE.open("r", encoding="utf-8") as f:
            history_messages = json.load(f)
        console.print("[green]历史聊天记录读取成功[/green]")
        return [
            item
            for item in history_messages
            if isinstance(item, dict) and item.get("role") in {"user", "assistant"}
        ]
    except Exception as exc:
        console.print(f"[red]聊天记录读取失败：{exc}[/red]")
        logging.exception("Failed to load chat history")
        return []


def trim_history(history: list[dict]) -> list[dict]:
    if len(history) <= MAX_HISTORY_MESSAGES:
        return history
    return history[-MAX_HISTORY_MESSAGES:]


def save_chat_history(history: list[dict]) -> None:
    with CHAT_FILE.open("w", encoding="utf-8") as f:
        json.dump(trim_history(history), f, ensure_ascii=False, indent=2)


history_messages = trim_history(load_chat_history())


class LLMClient:
    def __init__(self, config: dict):
        self.models = config["models"]
        self.retry_times = int(config["retry_times"])
        self.temperature = float(config["temperature"])
        self.client = OpenAI(
            api_key=os.getenv("API_KEY"),
            base_url=config["api_base_url"],
            timeout=float(config["timeout"]),
        )

    def chat_stream(self, messages: list[dict]) -> Iterable:
        last_error = None
        for model in self.models:
            for attempt in range(1, self.retry_times + 2):
                try:
                    logging.info("Requesting model=%s attempt=%s", model, attempt)
                    return self.client.chat.completions.create(
                        model=model,
                        messages=messages,
                        temperature=self.temperature,
                        stream=True,
                    )
                except Exception as exc:
                    last_error = exc
                    logging.exception("LLM request failed: model=%s attempt=%s", model, attempt)
                    time.sleep(min(2 * attempt, 5))

        raise RuntimeError(f"所有模型请求均失败：{last_error}")


llm_client = LLMClient(CONFIG)


def build_messages(user_input: str, history: list[dict]) -> list[dict]:
    return [
        {"role": "system", "content": build_system_prompt(user_input)},
        *trim_history(history),
        {"role": "user", "content": user_input},
    ]


def multiline_input() -> str:
    kb = KeyBindings()

    @kb.add("enter")
    def _(event):
        event.app.exit(result=event.app.current_buffer.text)

  

    @kb.add("c-j")
    def _(event):
        event.app.current_buffer.insert_text("\n")

    @kb.add("escape", "j")
    def _(event):
        event.app.current_buffer.insert_text("\n")

    session = PromptSession(multiline=True, key_bindings=kb)
    return session.prompt(
        "Sean: ",
        bottom_toolbar="Enter 发送；alt+j  换行；/upload \"文件路径\" 上传文件；输入 exit 退出",
    )


class SpaceStopper:
    """监听空格中断 AI 流式输出。"""

    def __init__(self):
        self.stop_event = threading.Event()
        self.thread = None
        self.running = False
        self.old_terminal_settings = None

    def start(self) -> None:
        self.running = True
        self.stop_event.clear()
        self.thread = threading.Thread(target=self._listen, daemon=True)
        self.thread.start()

    def stop(self) -> None:
        self.running = False
        if self.thread is not None:
            self.thread.join(timeout=0.2)
        self._restore_terminal()

    def stopped(self) -> bool:
        return self.stop_event.is_set()

    def _listen(self) -> None:
        if os.name == "nt":
            self._listen_windows()
        else:
            self._listen_unix()

    def _listen_windows(self) -> None:
        import msvcrt

        while self.running:
            if msvcrt.kbhit() and msvcrt.getwch() == " ":
                self.stop_event.set()
                self.running = False
                break
            time.sleep(0.03)

    def _listen_unix(self) -> None:
        import select
        import termios
        import tty

        try:
            self.old_terminal_settings = termios.tcgetattr(sys.stdin)
            tty.setcbreak(sys.stdin.fileno())
            while self.running:
                readable, _, _ = select.select([sys.stdin], [], [], 0.03)
                if readable and sys.stdin.read(1) == " ":
                    self.stop_event.set()
                    self.running = False
                    break
        except Exception:
            logging.exception("Space stopper failed")
        finally:
            self._restore_terminal()

    def _restore_terminal(self) -> None:
        if os.name == "nt" or self.old_terminal_settings is None:
            return
        try:
            import termios

            termios.tcsetattr(sys.stdin, termios.TCSADRAIN, self.old_terminal_settings)
            self.old_terminal_settings = None
        except Exception:
            logging.exception("Failed to restore terminal settings")


def render_formatted_reply(text: str) -> None:
    pattern = r"\$\$(.*?)\$\$"
    parts = re.split(pattern, text, flags=re.S)

    for i, part in enumerate(parts):
        if i % 2 == 0:
            if part.strip():
                console.print(Markdown(part))
        else:
            formula = part.strip()
            console.print(
                Panel(
                    Text(formula, style="bold magenta"),
                    title="LaTeX",
                    border_style="magenta",
                )
            )


def extract_code_blocks(reply: str) -> list[tuple[str, str]]:
    pattern = re.compile(r"```([^\n`]*)\n(.*?)```", flags=re.S)
    return [(match.group(1).strip(), match.group(2)) for match in pattern.finditer(reply)]


def markdown_to_copyable_html(markdown_text: str) -> str:
    code_pattern = re.compile(r"```([^\n`]*)\n(.*?)```", flags=re.S)
    chunks = []
    last_end = 0

    for index, match in enumerate(code_pattern.finditer(markdown_text), start=1):
        before = markdown_text[last_end:match.start()]
        if before.strip():
            chunks.append(f'<section class="markdown-text"><pre>{html.escape(before)}</pre></section>')

        info = html.escape(match.group(1).strip() or "code")
        code = html.escape(match.group(2).rstrip())
        chunks.append(
            f"""
            <section class="code-card">
                <div class="code-toolbar">
                    <span>{info}</span>
                    <button type="button" onclick="copyCode('code-{index}', this)">复制</button>
                </div>
                <pre><code id="code-{index}">{code}</code></pre>
            </section>
            """
        )
        last_end = match.end()

    tail = markdown_text[last_end:]
    if tail.strip():
        chunks.append(f'<section class="markdown-text"><pre>{html.escape(tail)}</pre></section>')

    return "\n".join(chunks)


def write_copyable_reply_html(reply: str) -> Path | None:
    if not extract_code_blocks(reply):
        return None

    body = markdown_to_copyable_html(reply)
    page = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>AI Reply</title>
  <style>
    :root {{
      color-scheme: light dark;
      --bg: #f6f7f9;
      --panel: #ffffff;
      --text: #1f2937;
      --muted: #667085;
      --border: #d0d5dd;
      --accent: #1769aa;
      --code-bg: #101828;
      --code-text: #e6edf3;
    }}
    body {{
      margin: 0;
      background: var(--bg);
      color: var(--text);
      font-family: "Segoe UI", Arial, sans-serif;
      line-height: 1.6;
    }}
    main {{
      width: min(980px, calc(100% - 32px));
      margin: 28px auto;
    }}
    h1 {{
      margin: 0 0 18px;
      font-size: 24px;
      font-weight: 650;
    }}
    .markdown-text,
    .code-card {{
      background: var(--panel);
      border: 1px solid var(--border);
      border-radius: 8px;
      margin: 14px 0;
      overflow: hidden;
    }}
    .markdown-text pre {{
      white-space: pre-wrap;
      margin: 0;
      padding: 16px;
      font: inherit;
    }}
    .code-toolbar {{
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 12px;
      padding: 10px 12px;
      border-bottom: 1px solid var(--border);
      color: var(--muted);
      font-size: 14px;
    }}
    button {{
      border: 1px solid var(--accent);
      background: var(--accent);
      color: white;
      border-radius: 6px;
      padding: 6px 12px;
      cursor: pointer;
      font-size: 14px;
    }}
    pre {{
      margin: 0;
      overflow: auto;
    }}
    code {{
      display: block;
      padding: 16px;
      background: var(--code-bg);
      color: var(--code-text);
      font: 14px/1.5 Consolas, "Cascadia Code", monospace;
      white-space: pre;
    }}
  </style>
</head>
<body>
  <main>
    <h1>AI 回复</h1>
    {body}
  </main>
  <script>
    async function copyCode(id, button) {{
      const text = document.getElementById(id).innerText;
      await navigator.clipboard.writeText(text);
      const oldText = button.innerText;
      button.innerText = "已复制";
      setTimeout(() => button.innerText = oldText, 1200);
    }}
  </script>
</body>
</html>
"""
    HTML_REPLY_FILE.write_text(page, encoding="utf-8")
    return HTML_REPLY_FILE


def extract_filenames_from_user_input(user_input: str) -> list[str]:
    pattern = r"([A-Za-z0-9_\-\u4e00-\u9fff]+?\.(?:py|cpp|md))"
    names = re.findall(pattern, user_input, flags=re.I)
    result = []
    for name in names:
        if name not in result:
            result.append(name)
    return result


def safe_filename(filename: str) -> str:
    filename = filename.strip().strip('"').strip("'")
    filename = filename.replace("\\", "/").split("/")[-1]
    filename = re.sub(r'[<>:"/\\|?*]', "_", filename)
    return filename or "generated.md"


def extract_filename_from_code_info(info: str) -> str | None:
    info = info.strip()
    filename_match = re.search(
        r'filename\s*=\s*["\']?([^"\'\s]+)["\']?',
        info,
        flags=re.I,
    )
    if filename_match:
        return safe_filename(filename_match.group(1))

    direct_match = re.search(
        r"([A-Za-z0-9_\-\u4e00-\u9fff]+?\.(?:py|cpp|md))",
        info,
        flags=re.I,
    )
    if direct_match:
        return safe_filename(direct_match.group(1))
    return None


def language_to_extension(language: str) -> str | None:
    language = language.lower().strip()
    if language in {"python", "py"}:
        return ".py"
    if language in {"cpp", "c++", "cplusplus"}:
        return ".cpp"
    if language in {"markdown", "md"}:
        return ".md"
    return None


def save_code_blocks(reply: str, user_input: str) -> list[Path]:
    matches = extract_code_blocks(reply)
    if not matches:
        return []

    requested_names = extract_filenames_from_user_input(user_input)
    requested_index = 0
    saved_files = []
    counters = {".py": 1, ".cpp": 1, ".md": 1}

    for info, code in matches:
        info_parts = info.split()
        language = info_parts[0].lower() if info_parts else ""
        extension = language_to_extension(language)
        if extension is None:
            continue

        filename = extract_filename_from_code_info(info)
        if filename is None and requested_index < len(requested_names):
            candidate = safe_filename(requested_names[requested_index])
            if candidate.lower().endswith(extension):
                filename = candidate
                requested_index += 1

        if filename is None:
            filename = f"generated_{counters[extension]}{extension}"
            counters[extension] += 1

        filename = safe_filename(filename)
        if not filename.lower().endswith(extension):
            filename += extension

        save_path = CURRENT_DIR / filename
        with save_path.open("w", encoding="utf-8", newline="\n") as f:
            f.write(code.rstrip() + "\n")
        saved_files.append(save_path)

    return saved_files


def parse_upload_command(user_input: str) -> tuple[Path, str] | None:
    text = user_input.strip()
    if not text.lower().startswith("/upload "):
        return None

    rest = text[len("/upload ") :].strip()
    quoted = re.match(r'^["\'](.+?)["\']\s*(.*)$', rest)
    if quoted:
        return resolve_path(quoted.group(1)), quoted.group(2).strip()
    return resolve_path(rest), ""


def build_uploaded_file_prompt(file_path: Path, question: str) -> str | None:
    if not file_path.exists():
        console.print(f"[red]没有找到文件：{file_path}[/red]")
        return None
    if not file_path.is_file():
        console.print(f"[red]这不是一个文件：{file_path}[/red]")
        return None
    try:
        suffix = file_path.suffix.lower()

        if suffix == ".pptx":
            content, truncated = load_ppt_file(
                file_path,
                int(CONFIG["max_upload_chars"]),
            )

        elif suffix == ".docx":
            content, truncated = load_docx_file(
                file_path,
                int(CONFIG["max_upload_chars"]),
            )

        elif suffix == ".pdf":
            content, truncated = load_pdf_file(
                file_path,
                int(CONFIG["max_upload_chars"]),
            )

        else:
            content, truncated = load_text_file(
                file_path,
                int(CONFIG["max_upload_chars"]),
            )
    
    except UnicodeDecodeError:
        console.print("[red]这个文件不是 UTF-8 文本文件，暂时不能上传二进制文件或非文本文件。[/red]")
        return None
    except Exception as exc:
        console.print(f"[red]文件读取失败：{exc}[/red]")
        return None

    question = question or "请阅读并分析这个文件，指出重点内容和可改进之处。"
    truncated_note = "\n\n[提示：文件内容较长，已截取前半部分。]" if truncated else ""
    console.print(f"[green]已上传文件：{file_path}[/green]")

    return f"""我上传了一个文件，请根据文件内容回答我的问题。

文件名：{file_path.name}
我的问题：{question}
{truncated_note}

文件内容：
```text
{content}
```"""


def preprocess_user_input(user_input: str) -> str | None:
    upload = parse_upload_command(user_input)
    if upload is None:
        return user_input
    return build_uploaded_file_prompt(*upload)


def stream_ai_reply(request_messages: list[dict], user_input: str):
    full_reply = ""
    cancelled = False
    start_time = time.time()

    console.print()
    console.print("[dim]提示：AI 思考时按空格可以中止。[/dim]")

    stopper = SpaceStopper()
    stopper.start()
    stream = None

    try:
        with console.status("[bold cyan]AI 正在检索知识、组织答案并流式生成...[/bold cyan]", spinner="dots12"):
            stream = llm_client.chat_stream(request_messages)

            for chunk in stream:
                if stopper.stopped():
                    cancelled = True
                    try:
                        stream.close()
                    except Exception:
                        logging.exception("Failed to close stream after cancellation")
                    break

                if not chunk.choices:
                    continue

                delta = chunk.choices[0].delta
                if delta.content is not None:
                    full_reply += delta.content
    except KeyboardInterrupt:
        cancelled = True
    finally:
        stopper.stop()
        if stream is not None:
            try:
                stream.close()
            except Exception:
                logging.exception("Failed to close stream")

    elapsed = time.time() - start_time
    console.print()
    console.print(f"[dim]思考耗时: {elapsed:.2f} 秒[/dim]")
    logging.info("AI reply completed cancelled=%s elapsed=%.2fs chars=%s", cancelled, elapsed, len(full_reply))

    if cancelled:
        console.print("[bold red]已中止 AI 输出。[/bold red]")
        console.print("[yellow]这次回复没有保存到聊天记录。[/yellow]")
        return full_reply, True, []

    console.print("\n[bold green]AI:[/bold green]\n")
    render_formatted_reply(full_reply)

    copyable_html = write_copyable_reply_html(full_reply)
    if copyable_html:
        console.print()
        console.print("[bold green]已生成带复制按钮的代码块预览：[/bold green]")
        console.print(f"[cyan]{copyable_html}[/cyan]")

    saved_files = save_code_blocks(full_reply, user_input)
    if saved_files:
        console.print()
        console.print("[bold green]已自动保存文件到当前目录：[/bold green]")
        for file_path in saved_files:
            console.print(f"[cyan]{file_path}[/cyan]")

    return full_reply, False, saved_files


def print_startup_help() -> None:
    console.print("[bold blue]=== My Assistant ===[/bold blue]")
    console.print("[dim]输入 exit 退出[/dim]")
    console.print("[dim]需要换行：Alt+j；[/dim]")
    console.print("[dim]AI 思考时按空格可以中止。[/dim]")
    console.print()
    console.print("[bold]上传文件：[/bold]")
    console.print('  /upload "D:\\Edge 下载\\d.md"')
    console.print('  /upload "D:\\Edge 下载\\d.md" 请总结这个文件')
    console.print("[dim]文件路径有空格时请加英文引号；也可以把文件拖到终端后补上 /upload。[/dim]")
    console.print("[dim]AI 回复里只要包含代码块，就会生成 latest_reply.html，打开后每个代码块都有复制按钮。[/dim]")
    console.print()


def main() -> None:
    global history_messages

    print_startup_help()

    while True:
        try:
            raw_input = multiline_input()
        except KeyboardInterrupt:
            console.print("\n[red]已退出。[/red]")
            break

        if raw_input.strip().lower() == "exit":
            break

        if not raw_input.strip():
            continue

        user_input = preprocess_user_input(raw_input)
        if user_input is None:
            continue

        request_messages = build_messages(user_input, history_messages)

        try:
            reply, cancelled, _ = stream_ai_reply(request_messages, user_input)
            if cancelled:
                continue

            history_messages.extend(
                [
                    {"role": "user", "content": user_input},
                    {"role": "assistant", "content": reply},
                ]
            )
            history_messages = trim_history(history_messages)
            save_chat_history(history_messages)
        except Exception as exc:
            logging.exception("Chat loop failed")
            console.print(f"\n[bold red]出错了：{exc}[/bold red]")


if __name__ == "__main__":
    main()
